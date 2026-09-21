"""Small, explicit industry mappings and parsers for public official datasets."""

from __future__ import annotations

import csv
import io
import math
import re
from datetime import datetime
from urllib.parse import urljoin, urlsplit

from parsel import Selector

from tradingagents.dataflows.fred import get_macro_data

from .transport import DocumentClient, SourceUnavailable, evidence

# Industry-specific statistics are enabled by business classification, never ticker guesses.
CATALOG = {
    "semiconductors": {"match": r"semiconductor|semiconductors", "fred": ["PCU334413334413", "INDPRO"],
                       "official": ["wsts"]},
    "energy": {"match": r"oil.*gas|petroleum|crude|refin(?:ing|eries)",
               "fred": ["DCOILWTICO", "DCOILBRENTEU"], "official": ["eia"]},
    "manufacturing": {"match": r"machinery|farm.*construction|agricultural.*construction|industrial.*manufactur|industrial.*equipment|metal.fabricat|automobile|auto.manufactur",
                      "fred": ["INDPRO", "PCU333120333120"], "official": ["census"]},
    "software": {"match": r"software|cloud|computer.programming", "fred": [], "official": []},
    "retail": {"match": r"retail|department.store", "fred": ["RSAFS", "UMCSENT"], "official": []},
    "housing": {"match": r"residential.construction|homebuild|building.material",
                "fred": ["HOUST", "PERMIT"], "official": []},
}


def select_industry(classification: str) -> list[str]:
    """Only select known mappings; absence is a coverage gap, not generic manufacturing."""
    return [name for name, spec in CATALOG.items() if re.search(spec["match"], classification, re.I)]


def parse_wsts(body: bytes, cutoff: str) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        if "Monthly Data" not in workbook.sheetnames:
            raise SourceUnavailable("WSTS Monthly Data sheet is missing")
        rows = list(workbook["Monthly Data"].values)
        units = " ".join(str(value) for row in rows[:4] for value in row if value is not None)
        if not re.search(r"1000\s*US\$", units, re.I):
            raise SourceUnavailable("WSTS units changed; cannot interpret values")
        records, year = [], None
        for row in rows:
            label = str(row[0] or "").strip()
            if re.fullmatch(r"(?:19|20)\d{2}", label):
                year = int(label)
            elif year and label in {"Americas", "Europe", "Japan", "Asia Pacific", "Worldwide"}:
                for month, value in enumerate(row[1:13], 1):
                    period = f"{year:04d}-{month:02d}"
                    if (period <= cutoff[:7] and year >= int(cutoff[:4]) - 2
                            and isinstance(value, (int, float)) and math.isfinite(value)):
                        records.append({"period": period, "region": label, "value": value,
                                        "display_usd_billion": round(value / 1_000_000, 3),
                                        "display_yi_usd": round(value / 100_000, 2)})
        if not records:
            raise SourceUnavailable("WSTS has no usable observations")
        periods = sorted({row["period"] for row in records})[-15:]
        return {"unit": "1000 USD", "frequency": "monthly", "seasonal_adjustment": "not specified",
                "latest_period": periods[-1], "observations": [row for row in records if row["period"] in periods]}
    finally:
        workbook.close()


def parse_census(body: bytes, label: str, cutoff: str, classification: str) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.values]
        first = next((i for i, row in enumerate(rows)
                      if isinstance(row[0], str) and sum(isinstance(v, (int, float)) for v in row[1:]) >= 3), None)
        if first is None:
            raise SourceUnavailable("Census table structure changed")
        headers = [[str(value) if value is not None else "" for value in row[:16]] for row in rows[:first]]
        heading = " ".join(cell for row in headers for cell in row)
        if "millions of dollars" not in heading.lower():
            raise SourceUnavailable("Census units changed")
        # Keep positional multirow headers intact, including p/r and adjustment labels.
        year_tokens = re.findall(r"\b(20\d{2})[pr]?\b", heading)
        if not year_tokens or max(year_tokens) > cutoff[:4]:
            raise SourceUnavailable("Census report period is unavailable or after the analysis date")
        column_period = None
        for column in range(2, min(5, len(headers[-1]))):
            tokens = [row[column] for row in headers if len(row) > column and row[column]]
            month = next((token for token in tokens if token in ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")), None)
            year = next((re.match(r"(20\d{2})[pr]?$", token) for token in tokens if re.match(r"(20\d{2})[pr]?$", token)), None)
            if month and year:
                column_period = datetime.strptime(f"{month} {year.group(1)}", "%B %Y").strftime("%Y-%m")
                break
        if not column_period or column_period > cutoff[:7]:
            raise SourceUnavailable("Census latest month is unavailable or after the analysis date")
        pattern = r"all manufacturing|machinery|farm|construction"
        if re.search(r"automobile|auto.manufactur", classification, re.I):
            pattern = r"all manufacturing|transportation|motor vehicle"
        elif re.search(r"metal", classification, re.I):
            pattern = r"all manufacturing|primary metals|fabricated metal"
        selected = []
        for number, row in enumerate(rows[first:], first + 1):
            if (isinstance(row[0], str) and re.search(pattern, row[0], re.I)
                    and sum(isinstance(v, (int, float)) for v in row[1:]) >= 3):
                selected.append({"excel_row": number, "values": row[:16]})
        footnotes = [str(row[0]) for row in rows[first:] if row[0]
                     and sum(isinstance(v, (int, float)) for v in row[1:]) == 0
                     and re.search(r"prelim|revis|semiconductor|seasonal|source|exclud", str(row[0]), re.I)]
        if not selected:
            raise SourceUnavailable("No matching Census industry rows")
        return {"table": label, "unit": "million USD; percent-change columns are percentages",
                "latest_period": column_period, "headers_by_column_position": headers,
                "rows": selected[:8], "footnotes": footnotes,
                "interpretation": "Blank header cells belong to merged groups; preserve SA/NSA and p/r labels. Industry groups are not company-specific orders."}
    finally:
        workbook.close()


def parse_eia(body: bytes, table: str, cutoff: str) -> dict:
    rows = list(csv.reader(io.StringIO(body.decode("cp1252"))))
    offset = 1 if table == "stocks" else 2
    if not rows or len(rows[0]) < offset + 2:
        raise SourceUnavailable("EIA CSV header is missing")
    try:
        periods = [datetime.strptime(value, "%m/%d/%y").date().isoformat()
                   for value in rows[0][offset:offset + 2]]
    except ValueError:
        raise SourceUnavailable("EIA weekly header changed") from None
    if any(stamp > cutoff for stamp in periods):
        raise SourceUnavailable("EIA observations are after the analysis date")
    selected = []
    wanted = ({"Crude Oil", "Commercial (Excluding SPR)", "Strategic Petroleum Reserve (SPR)",
               "Total Motor Gasoline", "Distillate Fuel Oil"} if table == "stocks" else
              {"Crude Oil Inputs", "Gross Inputs", "Operable Capacity", "Percent Utilization"})
    for row in rows[1:]:
        if len(row) < offset + 2 or row[offset - 1].strip() not in wanted:
            continue
        label = row[offset - 1].strip()
        unit = "million barrels" if table == "stocks" else (
            "percent" if "Utilization" in label else "thousand barrels per day")
        try:
            values = [float(value.replace(",", "")) for value in row[offset:offset + 2]]
        except ValueError:
            continue
        selected.append({"metric": label, "unit": unit,
                         "observations": [{"period": period, "value": value}
                                          for period, value in zip(periods, values, strict=True)]})
    if not selected:
        raise SourceUnavailable("No matching EIA rows")
    return {"table": table, "frequency": "weekly", "latest_period": periods[0], "metrics": selected}


def collect_indicators(client: DocumentClient, classification: str, cutoff: str,
                       sources: set[str], *, live: bool) -> dict:
    families = select_industry(classification)
    records, failures = [], []
    series = list(dict.fromkeys(series for family in families for series in CATALOG[family]["fred"]))
    # Construction machinery PPI is not a valid proxy for all manufacturers.
    if not re.search(r"farm|construction|machinery", classification, re.I):
        series = [series_id for series_id in series if series_id != "PCU333120333120"]
    if "fred" in sources:
        for series_id in series[:4]:
            try:
                data = get_macro_data(series_id, cutoff, 730)
                if "No observations" in data or "not found" in data or data.startswith("FRED:"):
                    raise SourceUnavailable("FRED series has no evidence at this date")
                records.append(evidence("fred", f"https://fred.stlouisfed.org/series/{series_id}", data,
                                        period=f"historical vintage <= {cutoff}", notes="FRED vintage-pinned; see observation units and seasonal adjustment."))
            except Exception as exc:
                failures.append({"source": "fred", "series": series_id, "status": "unavailable", "reason": type(exc).__name__})
    official = {source for family in families for source in CATALOG[family]["official"]} & sources
    for source in sorted(official):
        if not live:
            failures.append({"source": source, "status": "unavailable", "reason": "Current file has no verified historical vintage"})
            continue
        try:
            if source == "wsts":
                index_url = "https://www.wsts.org/67/Historical-Billings-Report"
                index = client.fetch(index_url, source="wsts")
                selector = Selector(text=index["body"].decode("utf-8", errors="replace"))
                links = [urljoin(index_url, href) for href in selector.css("a::attr(href)").getall()
                         if ".xlsx" in href.lower()]
                url = next((link for link in links if urlsplit(link).hostname in {"www.wsts.org", "wsts.org"}), None)
                if not url:
                    raise SourceUnavailable("WSTS public XLSX link not found")
                doc = client.fetch(url, source="wsts")
                data = parse_wsts(doc["body"], cutoff)
                records.append(evidence(source, url, data, retrieved_at=doc["retrieved_at"],
                                        period=data["latest_period"], notes="Current public totals, potentially revised. No product mix, volume or ASP. Cite WSTS; do not reproduce the workbook."))
            elif source == "census":
                for table, label in ((1, "shipments"), (2, "new_orders"), (4, "inventories")):
                    url = f"https://www.census.gov/manufacturing/m3/prel/table{table}p.xlsx"
                    doc = client.fetch(url, source="census")
                    data = parse_census(doc["body"], label, cutoff, classification)
                    records.append(evidence(source, url, data, retrieved_at=doc["retrieved_at"], period=data["latest_period"],
                                            notes="Current revised vintage. Census new orders exclude semiconductor new orders."))
            else:
                for number, label in ((1, "stocks"), (2, "refining")):
                    url = f"https://ir.eia.gov/wpsr/table{number}.csv"
                    doc = client.fetch(url, source="eia")
                    data = parse_eia(doc["body"], label, cutoff)
                    records.append(evidence(source, url, data, retrieved_at=doc["retrieved_at"], period=data["latest_period"],
                                            notes="Current EIA weekly file; retrieval date is not publication date. Prior observations may be revised."))
        except Exception as exc:
            failures.append({"source": source, "status": "unavailable", "reason": str(exc) if isinstance(exc, SourceUnavailable) else type(exc).__name__})
    return {"families": families, "classification": classification, "evidence": records, "failures": failures,
            "coverage_note": "Only explicitly mapped indicators are included. Missing statistics do not imply weak demand or a negative outlook."}
